"""Preprocessing and leakage-safe splitting for dialogue MADRS data."""

from __future__ import annotations

import csv
import json
import random
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class DialogueSample:
    """One symptom label paired with patient speech from one session/item."""

    session_id: str
    symptom: str
    patient_text: str
    severity_score: float
    # Optional but strongly preferred: splitting uses this instead of session_id.
    patient_id: str | None = None

    def __post_init__(self) -> None:
        if not self.session_id.strip():
            raise ValueError("session_id must not be empty")
        if not self.symptom.strip():
            raise ValueError("symptom must not be empty")
        if not self.patient_text.strip():
            raise ValueError("patient_text must not be empty")
        if not 0.0 <= float(self.severity_score) <= 6.0:
            raise ValueError(
                f"MADRS severity_score must be in [0, 6], got {self.severity_score}"
            )


def format_symptom_input(
    symptom: str, patient_text: str, sep_token: str = "[SEP]"
) -> str:
    """Create the only model input used by the dialogue experiment.

    No therapist/interviewer text is accepted by this function, making the
    patient-only constraint explicit at the model boundary.
    """

    symptom = " ".join(str(symptom).split())
    patient_text = str(patient_text).strip()
    if not symptom or not patient_text:
        raise ValueError("symptom and patient_text must both be non-empty")
    return f"SYMPTOM: {symptom} {sep_token} {patient_text}"


def _read_rows(path: str | Path) -> list[dict[str, Any]]:
    path = Path(path)
    if path.suffix.casefold() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [dict(row) for row in csv.DictReader(handle)]
    if path.suffix.casefold() in {".jsonl", ".ndjson"}:
        rows = []
        with path.open("r", encoding="utf-8") as handle:
            for line_number, line in enumerate(handle, 1):
                if not line.strip():
                    continue
                value = json.loads(line)
                if not isinstance(value, dict):
                    # This is malformed file content, not a caller type error.
                    raise ValueError(  # noqa: TRY004
                        f"Expected object at {path}:{line_number}"
                    )
                rows.append(value)
        return rows
    raise ValueError(f"Expected a .csv or .jsonl file: {path}")


def _patient_utterances(row: Mapping[str, Any]) -> list[str]:
    """Extract patient speech from supported transcript row shapes only."""

    # dyadic-sim transcript.jsonl: every row has separate therapist/patient fields.
    if "patient_text" in row:
        text = str(row.get("patient_text") or "").strip()
        return [text] if text else []

    # Common one-utterance-per-row transcript shape.
    speaker = str(row.get("speaker") or row.get("role") or "").casefold().strip()
    if speaker in {"patient", "client", "participant", "subject"}:
        text = str(row.get("text") or row.get("utterance") or "").strip()
        return [text] if text else []

    # A session object may embed turns/messages.
    turns = row.get("turns") or row.get("dialogue") or row.get("messages")
    if isinstance(turns, list):
        output: list[str] = []
        for turn in turns:
            if isinstance(turn, Mapping):
                output.extend(_patient_utterances(turn))
        return output
    return []


def extract_patient_text(
    transcript_path: str | Path, turn_separator: str = "\n"
) -> str:
    """Read a transcript and concatenate patient turns in chronological order."""

    utterances: list[str] = []
    for row in _read_rows(transcript_path):
        utterances.extend(_patient_utterances(row))
    if not utterances:
        raise ValueError(f"No patient utterances found in {transcript_path}")
    return turn_separator.join(utterances)


def build_dialogue_samples(
    transcripts_root: str | Path,
    labels_path: str | Path,
    *,
    transcript_name: str = "transcript.jsonl",
) -> list[DialogueSample]:
    """Join session/item labels to transcripts and keep patient turns only.

    ``labels_path`` must contain ``session_id``, ``symptom``, and
    ``severity_score``. ``patient_id`` and ``transcript_path`` are optional.
    When ``transcript_path`` is absent, the path is resolved as
    ``transcripts_root/session_id/transcript_name``.
    """

    root = Path(transcripts_root)
    text_cache: dict[Path, str] = {}
    samples: list[DialogueSample] = []
    seen: set[tuple[str, str]] = set()
    for row_number, row in enumerate(_read_rows(labels_path), 2):
        missing = [
            key
            for key in ("session_id", "symptom", "severity_score")
            if row.get(key) in (None, "")
        ]
        if missing:
            raise ValueError(f"Label row {row_number} is missing: {', '.join(missing)}")
        session_id = str(row["session_id"]).strip()
        symptom = str(row["symptom"]).strip()
        key = (session_id, symptom.casefold())
        if key in seen:
            raise ValueError(f"Duplicate label for session/symptom: {key}")
        seen.add(key)

        supplied_path = str(row.get("transcript_path") or "").strip()
        transcript_path = (
            Path(supplied_path)
            if supplied_path
            else root / session_id / transcript_name
        )
        if not transcript_path.is_absolute():
            transcript_path = root / transcript_path
        transcript_path = transcript_path.resolve()
        if transcript_path not in text_cache:
            text_cache[transcript_path] = extract_patient_text(transcript_path)

        patient_id = str(row.get("patient_id") or "").strip() or None
        samples.append(
            DialogueSample(
                session_id=session_id,
                symptom=symptom,
                patient_text=text_cache[transcript_path],
                severity_score=float(row["severity_score"]),
                patient_id=patient_id,
            )
        )
    if not samples:
        raise ValueError("No labelled samples were produced")
    return samples


def save_samples(samples: Iterable[DialogueSample], path: str | Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        for sample in samples:
            handle.write(json.dumps(asdict(sample), ensure_ascii=False) + "\n")


def load_samples(path: str | Path) -> list[DialogueSample]:
    return [
        DialogueSample(
            session_id=str(row["session_id"]),
            symptom=str(row["symptom"]),
            patient_text=str(row["patient_text"]),
            severity_score=float(row["severity_score"]),
            patient_id=(str(row["patient_id"]) if row.get("patient_id") else None),
        )
        for row in _read_rows(path)
    ]


def build_synthetic_workbook_samples(path: str | Path) -> list[DialogueSample]:
    """Read SyntheticData.xlsx as one MADRS item sample per subject/topic.

    The workbook repeats the same item score on several patient utterances.
    We concatenate those utterances in workbook order and use that score once.
    ``SPEAKER_00`` interviewer rows and the non-MADRS ``Allgemein`` topic are
    never included in model inputs or labels.
    """

    from openpyxl import load_workbook

    from dyadic_sim.scoring.instruments.madrs import TOPICS

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Expected one worksheet in the synthetic MADRS workbook")
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        expected = ("Subject", "Speaker", "Transcription", "Topic", "Score")
        if header is None or tuple(header[:5]) != expected:
            raise ValueError(f"Expected workbook columns {expected}")

        utterances: dict[tuple[str, str], list[str]] = {}
        scores: dict[tuple[str, str], float] = {}
        for row_number, row in enumerate(rows, start=2):
            subject, speaker, transcription, topic, score = row[:5]
            if speaker != "SPEAKER_01" or topic == "Allgemein":
                continue
            if topic not in TOPICS:
                raise ValueError(
                    f"Unknown MADRS topic on workbook row {row_number}: {topic!r}"
                )
            if not subject or not str(transcription or "").strip():
                raise ValueError(
                    f"Missing subject or patient speech on workbook row {row_number}"
                )
            if not isinstance(score, (int, float)) or isinstance(score, bool):
                # The spreadsheet contains an invalid label, not a bad API argument.
                raise ValueError(  # noqa: TRY004
                    f"Missing numeric score on workbook row {row_number}"
                )
            key = (str(subject), str(topic))
            numeric_score = float(score)
            if key in scores and scores[key] != numeric_score:
                raise ValueError(
                    f"Conflicting item scores for {key} on workbook row {row_number}"
                )
            scores[key] = numeric_score
            utterances.setdefault(key, []).append(str(transcription).strip())

        if not utterances:
            raise ValueError("No scored patient MADRS utterances found in workbook")
        return [
            DialogueSample(
                session_id=subject,
                patient_id=subject,
                symptom=topic,
                patient_text="\n".join(utterances[(subject, topic)]),
                severity_score=score,
            )
            for (subject, topic), score in scores.items()
        ]
    finally:
        workbook.close()


def grouped_split(
    samples: Sequence[DialogueSample],
    *,
    train_fraction: float = 0.70,
    validation_fraction: float = 0.15,
    seed: int = 42,
) -> dict[str, list[DialogueSample]]:
    """Split whole patients (or whole sessions) between train/val/test."""

    if not 0 < train_fraction < 1 or not 0 <= validation_fraction < 1:
        raise ValueError("Split fractions must be between zero and one")
    if train_fraction + validation_fraction >= 1:
        raise ValueError("train_fraction + validation_fraction must be less than one")
    if not samples:
        raise ValueError("Cannot split an empty sample collection")

    # If any patient IDs are supplied, require all of them: silently mixing
    # patient and session grouping could leak one partially identified patient.
    has_patient_ids = [sample.patient_id is not None for sample in samples]
    if any(has_patient_ids) and not all(has_patient_ids):
        raise ValueError("patient_id must be populated for every sample or none")
    group_for = (
        (lambda sample: f"patient:{sample.patient_id}")
        if all(has_patient_ids)
        else (lambda sample: f"session:{sample.session_id}")
    )
    groups = sorted({group_for(sample) for sample in samples})
    if len(groups) < 3:
        raise ValueError(
            "At least three patient/session groups are required for train/val/test"
        )
    random.Random(seed).shuffle(groups)

    train_count = max(1, round(len(groups) * train_fraction))
    val_count = max(1, round(len(groups) * validation_fraction))
    if train_count + val_count >= len(groups):
        train_count = len(groups) - 2
        val_count = 1
    train_groups = set(groups[:train_count])
    validation_groups = set(groups[train_count : train_count + val_count])
    output = {"train": [], "validation": [], "test": []}
    for sample in samples:
        group = group_for(sample)
        if group in train_groups:
            split = "train"
        elif group in validation_groups:
            split = "validation"
        else:
            split = "test"
        output[split].append(sample)
    return output
